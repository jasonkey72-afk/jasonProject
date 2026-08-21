"""
Office 대용량 파일 분할 엔진
=============================
xlsx / xlsm / docx / pptx / pdf 파일을 사용자가 지정한 용량(예: 10MB) 이하의
여러 개 파일로 나눈다.

분할 단위
    - xlsx/xlsm : 시트 → (한 시트가 너무 크면) 행 단위. 머리글 행은 매 파일마다 반복.
    - docx      : 문단/표 단위 (문서 흐름 순서 유지)
    - pptx      : 슬라이드 단위 (레이아웃/마스터/테마 그대로 유지)
    - pdf       : 쪽 단위 (글꼴·이미지·목차 그대로 유지)

동작 방식
    "지정한 용량 이하"를 보장하기 위해, 예상 크기로 한 번 저장해 본 뒤
    실제 파일 크기를 재서 넘치면 줄이고 너무 작으면 늘리는 방식으로 맞춘다.
    (압축률을 미리 알 수 없으므로 실제 저장 결과로 확인하는 것이 가장 정확하다)

원본 파일은 읽기만 하며 절대 수정하지 않는다.
이 모듈은 GUI와 분리되어 있어 단독으로 테스트할 수 있다.
"""

from __future__ import annotations

import os
import shutil
import threading
from copy import copy
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, List, Optional, Sequence, Tuple

KB = 1024
MB = 1024 * 1024

# 분할할 수 있는 확장자
SUPPORTED_EXTS = (".xlsx", ".xlsm", ".docx", ".pptx", ".pdf")
# 옛날 형식(분할 불가) — 새 형식으로 다시 저장해야 한다
LEGACY_EXTS = (".xls", ".doc", ".ppt")

# 저장 결과가 목표치보다 커지지 않도록 약간의 여유를 두고 채운다
FILL_RATIO = 0.93
# 이 비율보다 헐겁게 찼고 남은 내용이 있으면 더 채워 본다 (파일 개수를 줄이기 위함)
GROW_RATIO = 0.65
# 한 파트를 맞추기 위한 최대 저장 시도 횟수 (무한 루프 방지)
MAX_ATTEMPTS = 10
# 이 크기 이상이면 자동으로 "빠른 모드"(서식 제외, 메모리 절약)로 엑셀을 읽는다
AUTO_FAST_MODE_BYTES = 15 * MB


class Cancelled(Exception):
    """사용자가 중지를 눌렀을 때 발생."""


class SplitError(Exception):
    """분할할 수 없는 파일일 때 발생."""


def human_size(num_bytes: float) -> str:
    """1536000 → '1.5MB' 처럼 사람이 읽기 쉬운 문자열로."""
    num = float(num_bytes)
    for unit in ("B", "KB", "MB", "GB"):
        if abs(num) < 1024 or unit == "GB":
            if unit == "B":
                return f"{int(num)}{unit}"
            return f"{num:.1f}{unit}"
        num /= 1024
    return f"{num:.1f}GB"


@dataclass
class SplitOptions:
    """분할 옵션."""

    max_bytes: int = 10 * MB
    # None 이면 원본 파일과 같은 폴더에 저장
    output_dir: Optional[str] = None
    # 파일마다 "원본이름_분할" 하위 폴더를 만들어 그 안에 결과를 넣는다
    make_subfolder: bool = True
    # 엑셀: 각 파일 첫 줄에 머리글 행을 반복해서 넣는다
    repeat_header: bool = True
    header_rows: int = 1
    # 엑셀: 수식 대신 계산된 값을 넣는다 (분할하면 시트 간 참조가 깨지므로 기본 켬)
    formulas_to_values: bool = True
    # "auto" | "keep"(서식 유지) | "fast"(값만, 빠르고 메모리 절약)
    excel_mode: str = "auto"

    def resolve_output_dir(self, src: Path) -> Path:
        base = Path(self.output_dir) if self.output_dir else src.parent
        if self.make_subfolder:
            base = base / f"{src.stem}_분할"
        return base


@dataclass
class SplitResult:
    """파일 1개에 대한 분할 결과."""

    source: Path
    parts: List[Path] = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)
    output_dir: Optional[Path] = None

    @property
    def total_bytes(self) -> int:
        return sum(p.stat().st_size for p in self.parts if p.exists())


class Reporter:
    """진행 상황 알림 + 중지 요청 확인 (GUI가 없어도 동작한다)."""

    def __init__(
        self,
        log: Optional[Callable[[str], None]] = None,
        progress: Optional[Callable[[int, int, str], None]] = None,
        cancel_event: Optional[threading.Event] = None,
    ):
        self._log = log
        self._progress = progress
        self._cancel_event = cancel_event

    def log(self, message: str) -> None:
        if self._log:
            self._log(message)

    def progress(self, current: int, total: int, text: str = "") -> None:
        if self._progress:
            self._progress(current, total, text)

    def check_cancel(self) -> None:
        if self._cancel_event is not None and self._cancel_event.is_set():
            raise Cancelled()


# ---------------------------------------------------------------------------
# 공통 유틸
# ---------------------------------------------------------------------------

def _unique_dir(path: Path) -> Path:
    """이미 있는 폴더면 뒤에 (2), (3) ... 을 붙여 새 폴더 경로를 만든다."""
    if not path.exists():
        return path
    n = 2
    while True:
        candidate = path.with_name(f"{path.name} ({n})")
        if not candidate.exists():
            return candidate
        n += 1


def _part_path(out_dir: Path, stem: str, index: int, ext: str) -> Path:
    return out_dir / f"{stem}_part{index:02d}{ext}"


def _unique_stem(out_dir: Path, stem: str, ext: str) -> str:
    """같은 폴더에 같은 이름의 결과가 이미 있으면 덮어쓰지 않고 (2), (3)을 붙인다."""
    if not _part_path(out_dir, stem, 1, ext).exists():
        return stem
    n = 2
    while _part_path(out_dir, f"{stem} ({n})", 1, ext).exists():
        n += 1
    return f"{stem} ({n})"


def _safe_write(write_fn: Callable[[Path], None], out_path: Path) -> int:
    """저장 중 오류가 나면 반쯤 만들어진 파일을 남기지 않는다. 저장된 크기를 돌려준다."""
    try:
        write_fn(out_path)
    except Exception:
        if out_path.exists():
            try:
                out_path.unlink()
            except OSError:
                pass
        raise
    return out_path.stat().st_size


def _pack_units(
    total_units: int,
    write_part: Callable[[int, int, Path], int],
    out_dir: Path,
    stem: str,
    ext: str,
    max_bytes: int,
    reporter: Reporter,
    est_per_unit: float,
    warnings: List[str],
    unit_label: str,
    start_index: int = 1,
) -> Tuple[List[Path], int]:
    """
    total_units 개의 단위(슬라이드/문단 등)를 max_bytes 이하 파일들로 나눈다.

    write_part(start, end, out_path) 는 [start, end) 구간을 out_path 로 저장하고
    저장된 바이트 수를 돌려주는 함수다.

    반환값: (만들어진 파일 목록, 다음에 쓸 파트 번호)
    """
    parts: List[Path] = []
    target = int(max_bytes * FILL_RATIO)
    start = 0
    index = start_index

    while start < total_units:
        reporter.check_cancel()
        remaining = total_units - start
        take = remaining if est_per_unit <= 0 else int(target // max(est_per_unit, 1))
        take = max(1, min(remaining, take))

        out_path = _part_path(out_dir, stem, index, ext)
        shrunk = False
        size = 0
        for attempt in range(MAX_ATTEMPTS):
            reporter.check_cancel()
            reporter.progress(start, total_units, f"{index}번째 파일 만드는 중...")
            size = write_part(start, start + take, out_path)

            if size > max_bytes:
                if take == 1:
                    warnings.append(
                        f"{unit_label} {start + 1}번 하나만으로도 "
                        f"{human_size(size)}라서 지정한 용량을 넘습니다. (그대로 저장함)"
                    )
                    break
                shrunk = True
                new_take = int(take * (target / size))
                take = max(1, new_take if new_take < take else take - 1)
                continue

            # 너무 헐겁게 찼고 아직 남은 내용이 있으면 조금 더 담아 본다
            if (
                not shrunk
                and take < remaining
                and size < max_bytes * GROW_RATIO
                and attempt < MAX_ATTEMPTS - 1
            ):
                grown = int(take * (target / max(size, 1)))
                new_take = min(remaining, max(take + 1, grown))
                if new_take > take:
                    take = new_take
                    continue
            break

        parts.append(out_path)
        reporter.log(
            f"  · {out_path.name}  ({unit_label} {start + 1}~{start + take}, {human_size(size)})"
        )
        est_per_unit = max(size / take, 1.0)
        start += take
        index += 1

    reporter.progress(total_units, total_units, "")
    return parts, index


# ---------------------------------------------------------------------------
# PPTX — 슬라이드 단위
# ---------------------------------------------------------------------------

_R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_R_ID = f"{{{_R_NS}}}id"


def _pptx_keep_slides(prs, keep_start: int, keep_end: int) -> None:
    """[keep_start, keep_end) 범위 밖의 슬라이드를 프레젠테이션에서 제거한다."""
    sld_id_lst = prs.slides._sldIdLst
    for i, sld_id in enumerate(list(sld_id_lst)):
        if keep_start <= i < keep_end:
            continue
        r_id = sld_id.get(_R_ID)
        sld_id_lst.remove(sld_id)
        try:
            prs.part.drop_rel(r_id)
        except KeyError:
            pass


def split_pptx(src: Path, out_dir: Path, opts: SplitOptions, reporter: Reporter,
               stem: str) -> SplitResult:
    from pptx import Presentation

    result = SplitResult(source=src, output_dir=out_dir)

    prs = Presentation(str(src))
    total = len(prs.slides)
    del prs

    if total == 0:
        raise SplitError("슬라이드가 없는 파일입니다.")
    if total == 1:
        raise SplitError("슬라이드가 1장뿐이라 더 나눌 수 없습니다.")

    reporter.log(f"  슬라이드 {total}장을 나눕니다.")

    def write_part(start: int, end: int, out_path: Path) -> int:
        def _write(path: Path) -> None:
            part_prs = Presentation(str(src))
            _pptx_keep_slides(part_prs, start, end)
            part_prs.save(str(path))

        return _safe_write(_write, out_path)

    est = src.stat().st_size / total
    result.parts, _ = _pack_units(
        total, write_part, out_dir, stem, ".pptx", opts.max_bytes,
        reporter, est, result.warnings, "슬라이드",
    )
    return result


# ---------------------------------------------------------------------------
# PDF — 쪽(page) 단위
# ---------------------------------------------------------------------------

def _open_pdf(src: Path):
    """PDF 를 연다. 암호가 걸려 있으면 빈 암호로 한 번 시도해 본다."""
    from pypdf import PdfReader

    reader = PdfReader(str(src))
    if reader.is_encrypted:
        try:
            opened = reader.decrypt("")
        except Exception:
            opened = 0
        if not opened:
            raise SplitError(
                "암호가 걸린 PDF 는 나눌 수 없습니다. "
                "암호를 푼 뒤(다른 이름으로 저장) 다시 시도해 주세요."
            )
    return reader


def split_pdf(src: Path, out_dir: Path, opts: SplitOptions, reporter: Reporter,
              stem: str) -> SplitResult:
    from pypdf import PdfWriter

    result = SplitResult(source=src, output_dir=out_dir)

    reader = _open_pdf(src)
    total = len(reader.pages)
    del reader

    if total == 0:
        raise SplitError("쪽이 없는 PDF 입니다.")
    if total == 1:
        raise SplitError("1쪽짜리라 더 나눌 수 없습니다.")

    reporter.log(f"  {total}쪽을 나눕니다.")

    def write_part(start: int, end: int, out_path: Path) -> int:
        def _write(path: Path) -> None:
            part_reader = _open_pdf(src)
            writer = PdfWriter()
            try:
                # append 는 해당 쪽에 걸린 목차(북마크)까지 함께 가져온다
                writer.append(part_reader, pages=(start, end))
            except Exception:
                for i in range(start, end):
                    writer.add_page(part_reader.pages[i])
            try:
                if part_reader.metadata:
                    writer.add_metadata(part_reader.metadata)
            except Exception:
                pass
            with open(path, "wb") as f:
                writer.write(f)
            writer.close()

        return _safe_write(_write, out_path)

    est = src.stat().st_size / total
    result.parts, _ = _pack_units(
        total, write_part, out_dir, stem, ".pdf", opts.max_bytes,
        reporter, est, result.warnings, "쪽",
    )
    return result


# ---------------------------------------------------------------------------
# DOCX — 문단/표 단위
# ---------------------------------------------------------------------------

def _docx_blocks(doc):
    """본문의 최상위 요소(문단·표 등) 목록. 마지막 sectPr(용지 설정)은 제외한다."""
    from docx.oxml.ns import qn

    sect_pr = qn("w:sectPr")
    return [el for el in doc.element.body.iterchildren() if el.tag != sect_pr]


def _docx_prune_images(doc) -> None:
    """본문에서 더 이상 쓰이지 않는 그림을 문서에서 떼어내 용량을 실제로 줄인다."""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    used = set()
    for el in doc.element.body.iter():
        for key, value in el.attrib.items():
            if key.startswith(f"{{{_R_NS}}}"):
                used.add(value)

    part = doc.part
    for r_id, rel in list(part.rels.items()):
        if r_id in used or rel.is_external:
            continue
        if rel.reltype in (RT.IMAGE, RT.CHART):
            try:
                part.drop_rel(r_id)
            except KeyError:
                pass


def split_docx(src: Path, out_dir: Path, opts: SplitOptions, reporter: Reporter,
               stem: str) -> SplitResult:
    from docx import Document

    result = SplitResult(source=src, output_dir=out_dir)

    doc = Document(str(src))
    total = len(_docx_blocks(doc))
    del doc

    if total <= 1:
        raise SplitError("문단이 1개 이하라 더 나눌 수 없습니다.")

    reporter.log(f"  문단/표 {total}개를 나눕니다.")

    def write_part(start: int, end: int, out_path: Path) -> int:
        def _write(path: Path) -> None:
            part_doc = Document(str(src))
            body = part_doc.element.body
            for i, el in enumerate(_docx_blocks(part_doc)):
                if not (start <= i < end):
                    body.remove(el)
            _docx_prune_images(part_doc)
            part_doc.save(str(path))

        return _safe_write(_write, out_path)

    est = src.stat().st_size / total
    result.parts, _ = _pack_units(
        total, write_part, out_dir, stem, ".docx", opts.max_bytes,
        reporter, est, result.warnings, "문단",
    )
    return result


# ---------------------------------------------------------------------------
# XLSX / XLSM — 시트 단위 → (큰 시트는) 행 단위
# ---------------------------------------------------------------------------

class _RowSource:
    """
    시트의 행을 순서대로 공급한다.
    peek()/consume() 구조라서, 크기를 다시 맞추려고 행 수를 줄여도
    이미 읽어 둔 행을 버리지 않는다. (대용량 파일에서 다시 읽지 않기 위함)
    """

    def __init__(self, ws, header_rows: int, styled: bool):
        self.ws = ws
        self.styled = styled
        self._iter = ws.iter_rows() if styled else ws.iter_rows(values_only=True)
        self.header: List[tuple] = []
        for _ in range(max(0, header_rows)):
            try:
                self.header.append(next(self._iter))
            except StopIteration:
                break
        self._pending: List[tuple] = []
        self._eof = False

    def peek(self, n: int) -> List[tuple]:
        while not self._eof and len(self._pending) < n:
            try:
                self._pending.append(next(self._iter))
            except StopIteration:
                self._eof = True
        return self._pending[:n]

    def consume(self, n: int) -> None:
        del self._pending[:n]

    @property
    def exhausted(self) -> bool:
        return self._eof and not self._pending


def _row_is_empty(row, styled: bool) -> bool:
    if styled:
        return all(c.value is None for c in row)
    return all(v is None for v in row)


def _copy_dimensions(src_ws, dst_ws) -> None:
    """열 너비 등 보기 좋은 설정을 그대로 옮긴다. (읽기 전용 모드에서는 건너뜀)"""
    try:
        for key, dim in src_ws.column_dimensions.items():
            new_dim = dst_ws.column_dimensions[key]
            if dim.width:
                new_dim.width = dim.width
            new_dim.hidden = dim.hidden
    except Exception:
        pass


def _copy_merges(src_ws, dst_ws, src_start: int, src_end: int, offset: int, header_rows: int) -> None:
    """분할 범위 안에 온전히 들어가는 병합 셀만 옮긴다."""
    try:
        ranges = list(src_ws.merged_cells.ranges)
    except Exception:
        return
    for mr in ranges:
        if header_rows and mr.max_row <= header_rows:
            dst_ws.merge_cells(
                start_row=mr.min_row, start_column=mr.min_col,
                end_row=mr.max_row, end_column=mr.max_col,
            )
        elif src_start <= mr.min_row and mr.max_row <= src_end:
            dst_ws.merge_cells(
                start_row=mr.min_row + offset, start_column=mr.min_col,
                end_row=mr.max_row + offset, end_column=mr.max_col,
            )


class _StyleCopier:
    """
    셀 서식을 다른 통합문서로 옮긴다.

    openpyxl 의 cell._style 은 "그 파일 안에서의 서식 번호"라서 다른 파일에
    그대로 붙이면 깨진다. 그래서 글꼴/채우기/테두리 같은 실제 서식 객체를 옮기되,
    같은 서식은 한 번만 만들어 재사용한다. (대량 셀에서도 느려지지 않도록)
    """

    def __init__(self):
        self._cache = {}

    def apply(self, src_cell, dst_cell) -> None:
        key = src_cell.style_id
        entry = self._cache.get(key)
        if entry is None:
            entry = (
                copy(src_cell.font),
                copy(src_cell.fill),
                copy(src_cell.border),
                copy(src_cell.alignment),
                copy(src_cell.protection),
                src_cell.number_format,
            )
            self._cache[key] = entry
        (dst_cell.font, dst_cell.fill, dst_cell.border,
         dst_cell.alignment, dst_cell.protection, dst_cell.number_format) = entry


@dataclass
class _SheetBlock:
    """출력 파일에 들어갈 시트 한 개 분량."""

    title: str
    header: List[tuple]
    rows: List[tuple]
    src_ws: object = None


def _write_excel_part(out_path: Path, blocks: Sequence[_SheetBlock], styled: bool) -> int:
    from openpyxl import Workbook

    def _write(path: Path) -> None:
        if styled:
            wb = Workbook()
            wb.remove(wb.active)
        else:
            wb = Workbook(write_only=True)

        styler = _StyleCopier()
        for block in blocks:
            ws = wb.create_sheet(title=block.title[:31])
            if styled:
                out_row = 1
                for row in list(block.header) + list(block.rows):
                    for col_idx, cell in enumerate(row, start=1):
                        if cell.value is None and not cell.has_style:
                            continue
                        new_cell = ws.cell(row=out_row, column=col_idx, value=cell.value)
                        if cell.has_style:
                            styler.apply(cell, new_cell)
                    out_row += 1
                if block.src_ws is not None:
                    _copy_dimensions(block.src_ws, ws)
                    if block.rows:
                        src_start = block.rows[0][0].row
                        src_end = block.rows[-1][0].row
                        offset = len(block.header) + 1 - src_start
                        _copy_merges(block.src_ws, ws, src_start, src_end, offset,
                                     len(block.header))
                if block.header:
                    ws.freeze_panes = ws.cell(row=len(block.header) + 1, column=1)
            else:
                for row in list(block.header) + list(block.rows):
                    ws.append(list(row))

        wb.save(str(path))
        wb.close()

    return _safe_write(_write, out_path)


def _sheet_cell_count(ws) -> int:
    try:
        rows = ws.max_row or 0
        cols = ws.max_column or 0
    except Exception:
        return 0
    return max(0, rows) * max(0, cols)


def split_excel(src: Path, out_dir: Path, opts: SplitOptions, reporter: Reporter,
                stem: str) -> SplitResult:
    from openpyxl import load_workbook

    result = SplitResult(source=src, output_dir=out_dir)
    src_size = src.stat().st_size

    if opts.excel_mode == "keep":
        styled = True
    elif opts.excel_mode == "fast":
        styled = False
    else:
        styled = src_size < AUTO_FAST_MODE_BYTES

    if not styled:
        reporter.log("  빠른 모드로 읽습니다. (용량이 커서 서식 없이 데이터만 옮깁니다)")

    if src.suffix.lower() == ".xlsm":
        result.warnings.append("매크로(VBA)는 분할된 파일에 포함되지 않습니다. 결과는 .xlsx 로 저장됩니다.")

    wb = load_workbook(
        str(src),
        data_only=opts.formulas_to_values,
        read_only=not styled,
        keep_links=False,
    )
    try:
        sheets = [(name, wb[name]) for name in wb.sheetnames]
        if not sheets:
            raise SplitError("시트가 없는 파일입니다.")

        total_cells = sum(_sheet_cell_count(ws) for _, ws in sheets) or 1
        header_rows = opts.header_rows if opts.repeat_header else 0
        part_index = 1
        parts: List[Path] = []

        # 시트별 예상 크기로 "혼자서도 큰 시트"와 "묶어도 되는 작은 시트"를 나눈다
        est_sizes = {
            name: src_size * (_sheet_cell_count(ws) / total_cells) for name, ws in sheets
        }

        pending_group: List[Tuple[str, object]] = []
        pending_est = 0.0

        def flush_group() -> None:
            """모아 둔 작은 시트들을 파일 1개로 저장한다. (넘치면 시트를 덜어낸다)"""
            nonlocal pending_group, pending_est, part_index
            while pending_group:
                reporter.check_cancel()
                take = len(pending_group)
                out_path = _part_path(out_dir, stem, part_index, ".xlsx")
                while True:
                    blocks = []
                    for name, ws in pending_group[:take]:
                        source = _RowSource(ws, 0, styled)
                        blocks.append(_SheetBlock(name, [], source.peek(10 ** 9),
                                                  ws if styled else None))
                    size = _write_excel_part(out_path, blocks, styled)
                    if size <= opts.max_bytes or take == 1:
                        break
                    take = max(1, take - 1)

                names = ", ".join(name for name, _ in pending_group[:take])
                reporter.log(f"  · {out_path.name}  (시트: {names}, {human_size(size)})")
                if size > opts.max_bytes:
                    # 시트 1개인데도 넘친다 → 행 단위로 다시 나눈다
                    out_path.unlink(missing_ok=True)
                    name, ws = pending_group[0]
                    new_parts, part_index = _split_sheet_by_rows(
                        src, out_dir, stem, name, ws, styled, header_rows, opts,
                        reporter, result.warnings, part_index,
                    )
                    parts.extend(new_parts)
                else:
                    parts.append(out_path)
                    part_index += 1
                pending_group = pending_group[take:]
            pending_est = 0.0

        for name, ws in sheets:
            reporter.check_cancel()
            est = est_sizes[name]
            if est > opts.max_bytes * FILL_RATIO:
                flush_group()
                reporter.log(f"  시트 '{name}' 은(는) 커서 행 단위로 나눕니다.")
                new_parts, part_index = _split_sheet_by_rows(
                    src, out_dir, stem, name, ws, styled, header_rows, opts,
                    reporter, result.warnings, part_index,
                )
                parts.extend(new_parts)
                continue

            if pending_group and pending_est + est > opts.max_bytes * FILL_RATIO:
                flush_group()
            pending_group.append((name, ws))
            pending_est += est

        flush_group()

        if len(parts) <= 1:
            for p in parts:
                p.unlink(missing_ok=True)
            raise SplitError(
                "지정한 용량보다 작아서 나눌 필요가 없거나, 더 나눌 수 없는 파일입니다."
            )

        result.parts = parts
        return result
    finally:
        try:
            wb.close()
        except Exception:
            pass


def _split_sheet_by_rows(
    src: Path,
    out_dir: Path,
    stem: str,
    sheet_name: str,
    ws,
    styled: bool,
    header_rows: int,
    opts: SplitOptions,
    reporter: Reporter,
    warnings: List[str],
    part_index: int,
) -> Tuple[List[Path], int]:
    """시트 하나를 행 단위로 잘라 여러 파일로 저장한다. 머리글 행은 매번 반복한다."""
    source = _RowSource(ws, header_rows, styled)
    parts: List[Path] = []
    target = int(opts.max_bytes * FILL_RATIO)

    try:
        est_rows = max(1, (ws.max_row or 1) - header_rows)
    except Exception:
        est_rows = 1
    est_per_row = max(src.stat().st_size / est_rows, 1.0)
    row_no = header_rows

    while not source.exhausted:
        reporter.check_cancel()
        take = max(1, int(target // est_per_row))
        rows = source.peek(take)
        if not rows:
            break
        take = len(rows)

        out_path = _part_path(out_dir, stem, part_index, ".xlsx")
        shrunk = False
        size = 0
        for attempt in range(MAX_ATTEMPTS):
            reporter.check_cancel()
            reporter.progress(row_no, max(est_rows, row_no + take), f"{sheet_name} {row_no}행...")
            block = _SheetBlock(sheet_name, source.header, rows, ws if styled else None)
            size = _write_excel_part(out_path, [block], styled)

            if size > opts.max_bytes:
                if take == 1:
                    warnings.append(
                        f"'{sheet_name}' 시트의 {row_no + 1}행 하나만으로도 지정한 용량을 넘습니다."
                    )
                    break
                shrunk = True
                new_take = int(take * (target / size))
                take = max(1, new_take if new_take < take else take - 1)
                rows = source.peek(take)
                take = len(rows)
                continue

            if not shrunk and size < opts.max_bytes * GROW_RATIO and attempt < MAX_ATTEMPTS - 1:
                grown = max(take + 1, int(take * (target / max(size, 1))))
                more = source.peek(grown)
                if len(more) > take:
                    rows = more
                    take = len(more)
                    continue
            break

        source.consume(take)
        parts.append(out_path)
        reporter.log(
            f"  · {out_path.name}  ({sheet_name} {row_no + 1}~{row_no + take}행, {human_size(size)})"
        )
        est_per_row = max(size / take, 1.0)
        row_no += take
        part_index += 1

    return parts, part_index


# ---------------------------------------------------------------------------
# 진입점
# ---------------------------------------------------------------------------

def split_file(src_path, opts: SplitOptions, reporter: Optional[Reporter] = None) -> SplitResult:
    """파일 1개를 지정한 용량 이하의 여러 파일로 나눈다."""
    reporter = reporter or Reporter()
    src = Path(src_path)

    if not src.exists():
        raise SplitError("파일을 찾을 수 없습니다.")

    ext = src.suffix.lower()
    if ext in LEGACY_EXTS:
        raise SplitError(
            f"옛날 형식({ext})은 지원하지 않습니다. "
            f"Office에서 '다른 이름으로 저장'하여 최신 형식으로 바꾼 뒤 사용해 주세요."
        )
    if ext not in SUPPORTED_EXTS:
        raise SplitError(f"지원하지 않는 형식입니다: {ext}")

    size = src.stat().st_size
    if size <= opts.max_bytes:
        raise SplitError(
            f"이미 지정한 용량({human_size(opts.max_bytes)}) 이하입니다. "
            f"({human_size(size)}) 나눌 필요가 없습니다."
        )

    out_dir = opts.resolve_output_dir(src)
    if opts.make_subfolder:
        out_dir = _unique_dir(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    reporter.log(f"[{src.name}] {human_size(size)} → {human_size(opts.max_bytes)} 이하로 나눕니다.")

    out_ext = ".xlsx" if ext in (".xlsx", ".xlsm") else ext
    stem = _unique_stem(out_dir, src.stem, out_ext)

    try:
        if ext in (".xlsx", ".xlsm"):
            result = split_excel(src, out_dir, opts, reporter, stem)
        elif ext == ".docx":
            result = split_docx(src, out_dir, opts, reporter, stem)
        elif ext == ".pdf":
            result = split_pdf(src, out_dir, opts, reporter, stem)
        else:
            result = split_pptx(src, out_dir, opts, reporter, stem)
    except Cancelled:
        # 중지한 경우 만들다 만 결과물은 지운다
        if opts.make_subfolder and out_dir.exists():
            shutil.rmtree(out_dir, ignore_errors=True)
        raise
    except SplitError:
        if opts.make_subfolder and out_dir.exists() and not any(out_dir.iterdir()):
            out_dir.rmdir()
        raise

    over = [p for p in result.parts if p.stat().st_size > opts.max_bytes]
    reporter.log(
        f"  완료: {len(result.parts)}개 파일 생성"
        + (f" (그 중 {len(over)}개는 용량 초과)" if over else "")
    )
    for w in result.warnings:
        reporter.log(f"  ! {w}")
    return result


def split_files(paths: Sequence[str], opts: SplitOptions,
                reporter: Optional[Reporter] = None) -> List[SplitResult]:
    """여러 파일을 차례대로 나눈다. 한 파일이 실패해도 나머지는 계속 진행한다."""
    reporter = reporter or Reporter()
    results: List[SplitResult] = []
    for path in paths:
        reporter.check_cancel()
        try:
            results.append(split_file(path, opts, reporter))
        except Cancelled:
            raise
        except SplitError as e:
            reporter.log(f"[{Path(path).name}] 건너뜀 — {e}")
        except Exception as e:  # 예기치 못한 오류도 다음 파일로 넘어간다
            reporter.log(f"[{Path(path).name}] 실패 — {e}")
    return results


def collect_office_files(folder: str) -> List[str]:
    """폴더 안(하위 폴더 포함)의 분할 가능한 파일 목록."""
    found: List[str] = []
    for root, _dirs, files in os.walk(folder):
        for name in files:
            if name.startswith("~$"):
                continue
            if name.lower().endswith(SUPPORTED_EXTS):
                found.append(os.path.join(root, name))
    return sorted(found)
