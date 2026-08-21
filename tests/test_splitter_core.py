"""
분할 엔진 테스트
=================
실제로 큰 xlsx / docx / pptx 파일을 만들어 나눈 뒤,
  1) 결과 파일들이 모두 지정한 용량 이하인지
  2) 내용이 빠짐없이 들어갔는지
  3) 원본이 그대로인지
를 확인한다.

실행:  python tests/test_splitter_core.py
"""

import hashlib
import random
import shutil
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from splitter_core import (  # noqa: E402
    MB,
    SplitError,
    SplitOptions,
    Reporter,
    collect_office_files,
    human_size,
    split_file,
)

WORDS = ["보고", "실적", "매출", "부서", "검토", "요청", "데이터", "분석", "결과", "확인"]


def _rand_text(n_words=12, rng=None):
    rng = rng or random
    return " ".join(rng.choice(WORDS) + str(rng.randint(100, 99999)) for _ in range(n_words))


def make_xlsx(path, rows=40000, cols=12, sheets=1):
    from openpyxl import Workbook

    rng = random.Random(1)
    wb = Workbook(write_only=True)
    for s in range(sheets):
        ws = wb.create_sheet(title=f"데이터{s + 1}")
        ws.append([f"열{c + 1}" for c in range(cols)])
        for r in range(rows):
            ws.append([f"{r}-{c}-{rng.randint(0, 10**6)}" for c in range(cols)])
    wb.save(str(path))
    wb.close()


def make_docx(path, paragraphs=6000):
    from docx import Document

    rng = random.Random(2)
    doc = Document()
    doc.add_heading("부서 보고서", 0)
    for i in range(paragraphs):
        doc.add_paragraph(f"{i:05d} " + _rand_text(30, rng))
    doc.save(str(path))


def make_pptx(path, slides=120):
    from pptx import Presentation
    from pptx.util import Inches

    rng = random.Random(3)
    prs = Presentation()
    layout = prs.slide_layouts[1]
    for i in range(slides):
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = f"{i + 1}번 슬라이드"
        body = slide.placeholders[1].text_frame
        body.text = _rand_text(20, rng)
        for _ in range(12):
            body.add_paragraph().text = _rand_text(20, rng)
        slide.shapes.add_textbox(Inches(0.2), Inches(6.5), Inches(9), Inches(0.6)).text_frame.text = (
            _rand_text(10, rng)
        )
    prs.save(str(path))


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def check(condition, message):
    if not condition:
        raise AssertionError(message)
    print(f"    OK  {message}")


def run_case(name, src, limit, verify):
    print(f"\n[{name}] 원본 {human_size(src.stat().st_size)}  →  한 파일당 {human_size(limit)} 이하")
    before = sha256(src)
    logs = []
    reporter = Reporter(log=logs.append)
    opts = SplitOptions(max_bytes=limit, make_subfolder=True)
    result = split_file(src, opts, reporter)

    for line in logs:
        print("   ", line)

    check(len(result.parts) > 1, f"{len(result.parts)}개로 나뉘었다")
    over = [p for p in result.parts if p.stat().st_size > limit]
    check(not over, "모든 결과 파일이 지정한 용량 이하다 "
                    f"(최대 {human_size(max(p.stat().st_size for p in result.parts))})")
    check(sha256(src) == before, "원본 파일이 변경되지 않았다")
    check(all(p.exists() and p.stat().st_size > 0 for p in result.parts), "빈 파일이 없다")
    verify(result)
    # 파일 개수가 이론상 최소치의 3배를 넘지 않는지 (지나치게 잘게 쪼개지지 않는지)
    ideal = max(1, src.stat().st_size // limit)
    check(len(result.parts) <= ideal * 3 + 1,
          f"파일 개수가 합리적이다 ({len(result.parts)}개, 이론상 최소 {ideal}개)")
    return result


def verify_xlsx(src, sheets, rows, cols):
    def _verify(result):
        from openpyxl import load_workbook

        total_data_rows = 0
        headers_ok = True
        first_values = []
        for part in result.parts:
            wb = load_workbook(str(part), read_only=True)
            for ws in wb.worksheets:
                rows_iter = ws.iter_rows(values_only=True)
                header = next(rows_iter)
                if header[0] != "열1":
                    headers_ok = False
                for row in rows_iter:
                    if any(v is not None for v in row):
                        total_data_rows += 1
                        first_values.append(row[0])
            wb.close()
        check(headers_ok, "모든 파일의 첫 줄에 머리글이 들어 있다")
        check(total_data_rows == sheets * rows,
              f"데이터 행이 하나도 빠지지 않았다 ({total_data_rows}행)")
        check(len(set(first_values)) == len(first_values), "행이 중복되지 않았다")

    return _verify


def verify_docx(paragraphs):
    def _verify(result):
        from docx import Document

        texts = []
        for part in result.parts:
            doc = Document(str(part))
            texts.extend(p.text for p in doc.paragraphs)
        numbered = [t for t in texts if t[:5].isdigit()]
        check(len(numbered) == paragraphs, f"문단이 하나도 빠지지 않았다 ({len(numbered)}개)")
        check([t[:5] for t in numbered] == sorted(t[:5] for t in numbered), "문단 순서가 유지됐다")

    return _verify


def verify_pptx(slides):
    def _verify(result):
        from pptx import Presentation

        titles = []
        for part in result.parts:
            prs = Presentation(str(part))
            for slide in prs.slides:
                if slide.shapes.title is not None:
                    titles.append(slide.shapes.title.text)
        check(len(titles) == slides, f"슬라이드가 하나도 빠지지 않았다 ({len(titles)}장)")
        check(titles == [f"{i + 1}번 슬라이드" for i in range(slides)], "슬라이드 순서가 유지됐다")

    return _verify


def make_styled_xlsx(path, rows=3000, cols=8):
    """서식(색·굵게·숫자형식·열너비·병합)이 들어간 엑셀. '서식 유지' 모드 확인용."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    rng = random.Random(4)
    wb = Workbook()
    ws = wb.active
    ws.title = "실적표"
    ws["A1"] = "2026년 부서 실적"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=cols)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.append([f"열{c + 1}" for c in range(cols)])
    for cell in ws[2]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4F46E5")
    for c in range(cols):
        ws.column_dimensions[chr(ord("A") + c)].width = 18
    for r in range(rows):
        ws.append([rng.randint(1000, 999999) * 1.5 for _ in range(cols)])
        for cell in ws[ws.max_row]:
            cell.number_format = "#,##0.00"
    wb.save(str(path))


def verify_styled_xlsx(rows, cols):
    def _verify(result):
        from openpyxl import load_workbook

        data_rows = 0
        for part in result.parts:
            wb = load_workbook(str(part))
            ws = wb.active
            head = ws.cell(row=1, column=1)
            assert head.value == "2026년 부서 실적", "머리글이 반복되지 않았다"
            assert ws.cell(row=2, column=1).font.bold, "머리글 굵게 서식이 사라졌다"
            assert "A1:H1" in [str(r) for r in ws.merged_cells.ranges], "제목 병합이 사라졌다"
            body = ws.cell(row=3, column=1)
            assert body.number_format == "#,##0.00", f"숫자 서식이 사라졌다: {body.number_format}"
            assert ws.column_dimensions["A"].width == 18, "열 너비가 사라졌다"
            for r in range(3, ws.max_row + 1):
                if ws.cell(row=r, column=1).value is not None:
                    data_rows += 1
            wb.close()
        # 원본 1행(제목) + 2행(열 이름)이 머리글로 반복되므로 데이터 행 수만 비교
        check(data_rows >= rows, f"서식 유지 모드에서 데이터가 보존됐다 ({data_rows}행)")

    return _verify


def make_docx_with_images(path, blocks=40, image_kb=120):
    """그림이 많이 들어간 워드. 안 쓰는 그림이 잘 제거되는지 확인용."""
    import io

    from docx import Document
    from docx.shared import Inches
    from PIL import Image

    rng = random.Random(5)
    doc = Document()
    for i in range(blocks):
        doc.add_paragraph(f"{i:05d} " + _rand_text(15, rng))
        # 압축이 잘 안 되도록 무작위 픽셀 이미지를 만든다
        side = int((image_kb * 1024 / 3) ** 0.5)
        img = Image.frombytes(
            "RGB", (side, side), bytes(rng.randrange(256) for _ in range(side * side * 3))
        )
        buf = io.BytesIO()
        img.save(buf, format="PNG")
        buf.seek(0)
        doc.add_picture(buf, width=Inches(2))
    doc.save(str(path))


def verify_docx_images(blocks):
    def _verify(result):
        from docx import Document

        total_images = 0
        texts = []
        for part in result.parts:
            doc = Document(str(part))
            texts.extend(p.text for p in doc.paragraphs if p.text[:5].isdigit())
            total_images += sum(
                1 for rel in doc.part.rels.values() if "image" in rel.reltype
            )
        check(len(texts) == blocks, f"문단이 모두 보존됐다 ({len(texts)}개)")
        check(total_images == blocks,
              f"그림이 중복 없이 각 파일에 한 번씩만 들어갔다 ({total_images}개)")

    return _verify


def main():
    tmp = Path(tempfile.mkdtemp(prefix="split_test_"))
    print(f"작업 폴더: {tmp}")
    try:
        # --- 엑셀: 한 시트가 매우 큰 경우 (행 단위 분할)
        xlsx = tmp / "대용량_실적.xlsx"
        make_xlsx(xlsx, rows=40000, cols=12, sheets=1)
        run_case("XLSX / 큰 시트 1개", xlsx, 1 * MB, verify_xlsx(xlsx, 1, 40000, 12))

        # --- 엑셀: 여러 시트
        xlsx2 = tmp / "여러시트.xlsx"
        make_xlsx(xlsx2, rows=8000, cols=10, sheets=4)
        run_case("XLSX / 시트 4개", xlsx2, 1 * MB, verify_xlsx(xlsx2, 4, 8000, 10))

        # --- 워드
        docx = tmp / "부서_보고서.docx"
        make_docx(docx, paragraphs=6000)
        run_case("DOCX", docx, 200 * 1024, verify_docx(6000))

        # --- 파워포인트
        pptx = tmp / "발표자료.pptx"
        make_pptx(pptx, slides=120)
        run_case("PPTX", pptx, 200 * 1024, verify_pptx(120))

        # --- 엑셀: 서식 유지 모드
        styled = tmp / "서식표.xlsx"
        make_styled_xlsx(styled, rows=3000, cols=8)
        print(f"\n[XLSX / 서식 유지] 원본 {human_size(styled.stat().st_size)}")
        logs = []
        opts = SplitOptions(max_bytes=120 * 1024, excel_mode="keep", header_rows=2)
        res = split_file(styled, opts, Reporter(log=logs.append))
        for line in logs:
            print("   ", line)
        check(len(res.parts) > 1, f"{len(res.parts)}개로 나뉘었다")
        check(all(p.stat().st_size <= 120 * 1024 for p in res.parts),
              "모든 결과 파일이 지정한 용량 이하다")
        verify_styled_xlsx(3000, 8)(res)

        # --- 워드: 그림이 많은 문서 (안 쓰는 그림 제거 확인)
        img_docx = tmp / "사진보고서.docx"
        make_docx_with_images(img_docx, blocks=40, image_kb=120)
        run_case("DOCX / 그림 포함", img_docx, 1 * MB, verify_docx_images(40))

        # --- 이미 작은 파일은 건너뛴다
        print("\n[작은 파일 처리]")
        small = tmp / "작은파일.docx"
        make_docx(small, paragraphs=5)
        try:
            split_file(small, SplitOptions(max_bytes=10 * MB))
            raise AssertionError("작은 파일인데 오류가 나지 않았다")
        except SplitError as e:
            check("나눌 필요가 없습니다" in str(e), f"안내 메시지를 준다: {e}")

        # --- 옛날 형식 안내
        legacy = tmp / "옛날파일.xls"
        legacy.write_bytes(b"dummy")
        try:
            split_file(legacy, SplitOptions(max_bytes=1 * MB))
            raise AssertionError("xls 인데 오류가 나지 않았다")
        except SplitError as e:
            check("최신 형식" in str(e), "옛날 형식은 안내 메시지를 준다")

        # --- 폴더 수집
        print("\n[폴더 검색]")
        found = collect_office_files(str(tmp))
        check(all(f.lower().endswith((".xlsx", ".docx", ".pptx")) for f in found),
              f"폴더에서 Office 파일만 찾아낸다 ({len(found)}개)")

        print("\n모든 테스트를 통과했습니다.")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    main()
